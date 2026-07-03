#!/usr/bin/env python3
"""Strip Word content controls from a .docx, leaving the bracket text as plain runs.

Usage:
    python3 scripts/strip_content_controls.py INPUT.docx [OUTPUT.docx]

If OUTPUT is omitted the result is written to INPUT_stripped.docx.
"""
import copy
import sys
from pathlib import Path
from docx import Document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SDT  = f"{{{W}}}sdt"
SDT_CONTENT = f"{{{W}}}sdtContent"


def _strip_sdts(element):
    """Recursively replace w:sdt elements with their w:sdtContent children."""
    changed = True
    while changed:
        changed = False
        for sdt in element.findall(f".//{SDT}"):
            parent = sdt.getparent()
            if parent is None:
                continue
            idx = list(parent).index(sdt)
            content = sdt.find(SDT_CONTENT)
            if content is not None:
                for i, child in enumerate(list(content)):
                    parent.insert(idx + i, copy.deepcopy(child))
            parent.remove(sdt)
            changed = True  # restart search — tree mutated


def strip_content_controls(src: Path, dst: Path):
    doc = Document(str(src))
    _strip_sdts(doc.element.body)
    doc.save(str(dst))


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_stripped" + src.suffix)

    if not src.exists():
        print(f"File not found: {src}")
        sys.exit(1)

    strip_content_controls(src, dst)
    print(f"Saved → {dst}")

    # Quick verification
    from src.reports.template_inference import extract_placeholders
    tokens = extract_placeholders(dst)
    print(f"Placeholders found: {len(tokens)}")
    for tok in tokens[:20]:
        print(f"  {tok.delimiter}  {tok.inner!r}")
    if len(tokens) > 20:
        print(f"  … and {len(tokens) - 20} more")


if __name__ == "__main__":
    main()
